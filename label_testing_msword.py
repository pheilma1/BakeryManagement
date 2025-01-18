# import docx
# from docx.shared import Pt
# from docx.enum.style import WD_STYLE_TYPE
#
# doc = docx.Document()
# print (doc)
#
# parag = doc.add_paragraph("Hello!")
#
# font_styles = doc.styles
# font_charstyle = font_styles.add_style('IngredientStyle', WD_STYLE_TYPE.CHARACTER)
# font_object = font_charstyle.font
# font_object.size = Pt(20)
# font_object.name = 'Calibri'
#
# parag.add_run("Spinach Feta Pastry", style='IngredientStyle').bold = True
# parag.add_run("Python", style='IngredientStyle').italic = True
# parag = doc.add_paragraph("Python Normal")
# parag = doc.add_paragraph("Python Heading 1")
# parag.style = 'Heading 1'
#
# # parag.add_run("Python Normal", style='Normal').italic = True
# doc.save("test.docx")
#
#

from docx import Document
from docx.shared import Inches

# ------- initial code -------

document = Document()

p = document.add_paragraph()
r = p.add_run()
r.add_text('Good Morning every body,This is my ')
# picPath = 'D:/Development/Python/aa.png'
r.add_picture('LOGO.jpg')
r.add_text(' do you like it?')

document.save('demo.docx')

# ------- improved code -------

document = Document()

p = document.add_paragraph('Picture bullet section', 'List Bullet')
p = p.insert_paragraph_before('')
r = p.add_run()
r.add_picture('LOGO.jpg', width=Inches(1.0), height=Inches(.7))
p = p.insert_paragraph_before('My picture title', 'Heading 1')

document.save('demo_better.docx')